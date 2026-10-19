import base64
from collections import defaultdict
from datetime import datetime
from io import BytesIO

from odoo import _, fields, models
from odoo.exceptions import UserError


class InvoiceExcelImportWizard(models.TransientModel):
    _name = "invoice.excel.import.wizard"
    _description = "Importador flexible de facturas desde Excel"

    file_data = fields.Binary(required=True)
    file_name = fields.Char(required=True)
    template_id = fields.Many2one("account.invoice.import.template", required=True)
    company_id = fields.Many2one("res.company", default=lambda self: self.env.company, required=True)

    def action_import(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Debe cargar un archivo Excel."))

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UserError(_("Falta la libreria openpyxl. Instalela en el entorno de Odoo.")) from exc

        workbook = load_workbook(filename=BytesIO(base64.b64decode(self.file_data)), data_only=True)
        sheet = self._get_sheet(workbook)
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("El archivo no contiene datos."))

        header_index = self.template_id.header_row - 1
        if header_index >= len(rows):
            raise UserError(_("La fila de encabezado indicada no existe en el Excel."))

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_index]]
        mapping = self._build_mapping(headers)
        data_rows = rows[header_index + 1 :]

        grouped = defaultdict(list)
        for row in data_rows:
            row_data = self._row_to_canonical(headers, row, mapping)
            if not any(value not in (None, "") for value in row_data.values()):
                continue
            key = row_data.get("invoice_key") or "ROW-%s" % (len(grouped) + 1)
            grouped[key].append(row_data)

        created = self.env["account.move"]
        for invoice_rows in grouped.values():
            move = self._create_invoice(invoice_rows)
            created |= move

        return {
            "type": "ir.actions.act_window",
            "name": _("Facturas importadas"),
            "res_model": "account.move",
            "view_mode": "list,form",
            "domain": [("id", "in", created.ids)],
            "context": {"create": False},
        }

    def _get_sheet(self, workbook):
        if self.template_id.sheet_name:
            if self.template_id.sheet_name not in workbook.sheetnames:
                raise UserError(_("No existe la hoja '%s' en el archivo.") % self.template_id.sheet_name)
            return workbook[self.template_id.sheet_name]
        return workbook[workbook.sheetnames[0]]

    def _build_mapping(self, headers):
        mapping = {}
        template_lines = self.template_id.line_ids
        if not template_lines:
            raise UserError(_("La plantilla no tiene mapeos definidos."))

        headers_lc = {h.lower(): h for h in headers if h}
        for line in template_lines:
            target_header = headers_lc.get(line.excel_column.lower())
            if not target_header:
                raise UserError(
                    _("No se encontro la columna '%s' definida en la plantilla.") % line.excel_column
                )
            mapping[target_header] = line.target
        return mapping

    def _row_to_canonical(self, headers, row, mapping):
        result = {}
        for idx, header in enumerate(headers):
            if not header or header not in mapping:
                continue
            value = row[idx] if idx < len(row) else None
            result[mapping[header]] = value
        return result

    def _create_invoice(self, invoice_rows):
        header = invoice_rows[0]

        partner = self._resolve_partner(header)
        journal = self._resolve_journal(header)
        currency = self._resolve_currency(header)
        activity = self._resolve_activity(header)

        move_vals = {
            "move_type": (header.get("move_type") or "out_invoice").strip() if isinstance(header.get("move_type"), str) else (header.get("move_type") or "out_invoice"),
            "partner_id": partner.id,
            "invoice_date": self._parse_date(header.get("invoice_date")),
            "invoice_date_due": self._parse_date(header.get("invoice_date_due")),
            "ref": header.get("ref"),
            "journal_id": journal.id,
            "company_id": self.company_id.id,
            "activity_id": activity.id if activity else False,
            "currency_id": currency.id if currency else False,
        }

        line_vals = []
        for row in invoice_rows:
            line_vals.append((0, 0, self._prepare_line_vals(row, move_vals["move_type"])))

        move_vals["invoice_line_ids"] = line_vals
        move = self.env["account.move"].create(move_vals)
        return move

    def _resolve_partner(self, values):
        Partner = self.env["res.partner"].with_company(self.company_id)
        vat = values.get("partner_vat")
        name = values.get("partner_name")

        partner = False
        if vat:
            partner = Partner.search([("vat", "=", str(vat).strip())], limit=1)
        if not partner and name:
            partner = Partner.search([("name", "=ilike", str(name).strip())], limit=1)
        if not partner:
            if not name and not vat:
                raise UserError(_("No se pudo identificar cliente/proveedor: falta RUC o nombre."))
            partner = Partner.create({"name": name or vat, "vat": vat or False})
        return partner

    def _resolve_journal(self, values):
        Journal = self.env["account.journal"].with_company(self.company_id)
        raw = values.get("journal")
        if raw:
            raw = str(raw).strip()
            journal = Journal.search([("code", "=", raw)], limit=1)
            if not journal:
                journal = Journal.search([("name", "=ilike", raw)], limit=1)
            if journal:
                return journal

        move_type = values.get("move_type") or "out_invoice"
        if isinstance(move_type, str) and move_type.startswith("in_"):
            journal_type = "purchase"
        else:
            journal_type = "sale"
        journal = Journal.search([("type", "=", journal_type)], limit=1)
        if not journal:
            raise UserError(_("No existe un diario de tipo '%s' para la compania activa.") % journal_type)
        return journal

    def _resolve_currency(self, values):
        currency_code = values.get("currency")
        if not currency_code:
            return False
        return self.env["res.currency"].search([("name", "=", str(currency_code).strip().upper())], limit=1)

    def _resolve_activity(self, values):
        code = values.get("activity_code")
        name = values.get("activity_name")
        Activity = self.env["billing.activity"]
        activity = False
        if code:
            activity = Activity.search([("code", "=", str(code).strip())], limit=1)
        if not activity and name:
            activity = Activity.search([("name", "=ilike", str(name).strip())], limit=1)
        return activity

    def _prepare_line_vals(self, values, move_type):
        product = self._resolve_product(values.get("line_product"))
        account = self._resolve_account(values.get("line_account"), move_type, product)
        taxes = self._resolve_taxes(values.get("line_taxes"), move_type)

        return {
            "name": values.get("line_description") or (product.display_name if product else _("Linea importada")),
            "quantity": self._to_float(values.get("line_quantity"), default=1.0),
            "price_unit": self._to_float(values.get("line_price_unit"), default=0.0),
            "discount": self._to_float(values.get("line_discount"), default=0.0),
            "product_id": product.id if product else False,
            "account_id": account.id,
            "tax_ids": [(6, 0, taxes.ids)] if taxes else False,
        }

    def _resolve_product(self, raw):
        if not raw:
            return False
        Product = self.env["product.product"]
        raw = str(raw).strip()
        product = Product.search([("default_code", "=", raw)], limit=1)
        if not product:
            product = Product.search([("name", "=ilike", raw)], limit=1)
        return product

    def _resolve_account(self, raw, move_type, product=False):
        if raw:
            raw = str(raw).strip()
            account = self.env["account.account"].search([("code", "=", raw), ("company_id", "=", self.company_id.id)], limit=1)
            if not account:
                account = self.env["account.account"].search([("name", "=ilike", raw), ("company_id", "=", self.company_id.id)], limit=1)
            if account:
                return account

        if product:
            if isinstance(move_type, str) and move_type.startswith("in_"):
                account = product.property_account_expense_id or product.categ_id.property_account_expense_categ_id
            else:
                account = product.property_account_income_id or product.categ_id.property_account_income_categ_id
            if account:
                return account

        account_type = "expense" if isinstance(move_type, str) and move_type.startswith("in_") else "income"
        account = self.env["account.account"].search(
            [
                ("account_type", "=", account_type),
                ("deprecated", "=", False),
                ("company_id", "=", self.company_id.id),
            ],
            limit=1,
        )
        if not account:
            raise UserError(_("No se encontro cuenta contable para la linea importada."))
        return account

    def _resolve_taxes(self, raw, move_type):
        if not raw:
            return self.env["account.tax"]
        use_type = "purchase" if isinstance(move_type, str) and move_type.startswith("in_") else "sale"
        tax_codes = [item.strip() for item in str(raw).split(",") if item and item.strip()]
        taxes = self.env["account.tax"]
        for code in tax_codes:
            tax = self.env["account.tax"].search(
                [
                    ("type_tax_use", "=", use_type),
                    "|",
                    ("name", "=ilike", code),
                    ("description", "=ilike", code),
                ],
                limit=1,
            )
            if tax:
                taxes |= tax
        return taxes

    def _to_float(self, value, default=0.0):
        if value in (None, ""):
            return default
        if isinstance(value, (int, float)):
            return float(value)
        try:
            return float(str(value).replace(",", "."))
        except ValueError:
            return default

    def _parse_date(self, value):
        if not value:
            return False
        if isinstance(value, datetime):
            return fields.Date.to_date(value)
        if isinstance(value, str):
            value = value.strip()
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    continue
        return value
